"""Pitch tools (P2): screenshot config, deck builder, init-project. No browser, no network."""

import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import browser_steps  # noqa: E402
import init_project  # noqa: E402
import shots  # noqa: E402


def test_parse_size_and_output_name():
    assert shots.parse_size("1440x900") == (1440, 900)
    assert shots.output_name("home", "1440x900", ["1440x900"]) == "home.png"
    assert shots.output_name("home", "390x844", ["1440x900", "390x844"]) == "home@390x844.png"


def test_shots_config_rejects_duplicate_names(tmp_path):
    config = tmp_path / "shots.toml"
    config.write_text('[[shot]]\nname = "a"\n[[shot]]\nname = "a"\n')
    with pytest.raises(SystemExit, match="unique"):
        shots.load_shots(config)


def test_template_shots_config_is_valid():
    defaults, items = shots.load_shots(ROOT / "docs" / "pitch" / "shots.toml")
    assert defaults["sizes"] and {s["name"] for s in items} >= {"home", "feature-result"}


def test_unknown_step_is_reported_without_a_browser():
    with pytest.raises(browser_steps.StepError, match="no action"):
        browser_steps.run_step(object(), "http://x", {"text": "only data"})
    with pytest.raises(browser_steps.StepError, match="unknown action"):
        browser_steps.run_step(object(), "http://x", {"dance": "now"})
    # video-only steps are skipped when taking screenshots
    browser_steps.run_step(object(), "http://x", {"caption": "hi"})


def test_deck_builds_with_notes_and_warnings(tmp_path):
    pytest.importorskip("pptx")
    import deck

    (tmp_path / "deck.toml").write_text(
        """
[deck]
name = "t"
minutes = 0.5
[[slide]]
kind = "title"
title = "Name"
subtitle = "Tagline"
seconds = 10
notes = "Say hello."
[[slide]]
kind = "image"
title = "Product"
image = "shots/missing.png"
seconds = 30
[[slide]]
kind = "bullets"
title = "Too many words"
bullets = ["one two three four five six seven eight nine ten eleven twelve thirteen fourteen",
           "one two three four five six seven eight nine ten eleven twelve thirteen fourteen",
           "one two three four five six seven eight nine ten eleven twelve thirteen fourteen",
           "a", "b"]
seconds = 5
notes = "x"
"""
    )
    path, report = deck.build(tmp_path / "deck.toml", tmp_path / "out")
    from pptx import Presentation

    prs = Presentation(str(path))
    assert len(prs.slides) == 3
    assert prs.slides[0].notes_slide.notes_text_frame.text == "Say hello."
    text = "\n".join(report.warnings)
    assert "image not found" in text and "no speaker notes" in text
    assert "words on the slide" in text and "bullets" in text and "slot is 0.5 min" in text


def test_template_deck_has_no_warnings():
    pytest.importorskip("pptx")
    import deck

    _, report = deck.build(ROOT / "docs" / "pitch" / "deck.toml", Path("/tmp/hackkit-deck-test"))
    missing = [w for w in report.warnings if "image not found" not in w]
    assert missing == [] and report.seconds <= 240


def test_init_project_renames_everything(tmp_path, monkeypatch):
    for rel in ("web/config.json", "web/index.html", "docs/pitch/deck.toml",
                "docs/pitch/demo_flow.toml", "docs/SUBMISSION.md", "README.md"):  # fmt: skip
        target = tmp_path / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_text((ROOT / rel).read_text(encoding="utf-8"), encoding="utf-8")
    if not (tmp_path / "README.md").read_text().startswith("# hackkit"):
        (tmp_path / "README.md").write_text("# hackkit\n\ntemplate readme\n")
    monkeypatch.setattr(init_project, "ROOT", tmp_path)
    monkeypatch.setattr(init_project, "git_remote_pages_url", lambda: "https://me.github.io/app/")
    init_project.main(
        ["--name", "Acme Check", "--tagline", "Fast checks", "--team", "Team X · A, B"]
    )

    config = json.loads((tmp_path / "web/config.json").read_text())
    assert config["name"] == "Acme Check" and config["hero"]["title"] == "Fast checks"
    assert "<title>Acme Check</title>" in (tmp_path / "web/index.html").read_text()
    deck_text = (tmp_path / "docs/pitch/deck.toml").read_text()
    assert 'title = "Acme Check"' in deck_text and "Team X" in deck_text
    readme = (tmp_path / "README.md").read_text()
    assert readme.startswith("# Acme Check") and "https://me.github.io/app/" in readme
    assert (tmp_path / "docs/HACKKIT.md").read_text().startswith("# hackkit")
